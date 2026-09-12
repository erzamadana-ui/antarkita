# Membangun Matriks Uji (xlsx) + tabel Markdown dari satu sumber data.
# Jalankan: python3 docs/uji/buat-matriks.py
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TGL = "2026-09-12"
RC = "commit 30e3ced + perbaikan QC (0083/0084) · AAB v3.0.0 build 104 (aab-4) → build 105 setelah push"
PEL = "Claude (QA Lead independen, atas kuasa Erza)"

# (ID, bagian, skenario, prasyarat, data uji, langkah, expected, actual, severity, bukti, status, PIC/target)
T = [
 # 6.1 Akun & otorisasi
 ("AUTH-01","6.1 / E2E-09","Pengguna A tidak bisa membaca dompet/order pengguna B (IDOR)","Akun uji customer","2 akun uji","Sesi RLS sebagai customer → select wallets/orders milik orang lain","0 baris","0 baris","-","bukti/uji_keamanan.log","PASS",""),
 ("AUTH-02","6.1","Eskalasi hak: customer mengubah role sendiri jadi admin","Akun uji customer","update profiles.role","Jalankan update sebagai customer","Diblokir trigger","Diblokir","-","bukti/uji_keamanan.log","PASS",""),
 ("AUTH-03","6.1","Fungsi admin dipanggil non-admin / anon","-","admin_adjust_wallet, admin_map_status, admin_turn_status, admin_set_city_status…","Panggil RPC sebagai customer & anon","Ditolak 'Hanya admin'; anon tanpa hak EXECUTE","Ditolak","-","bukti/uji_keamanan.log, uji_kota.log 6a–6f, uji_turn.log 2a–3c","PASS",""),
 ("AUTH-04","6.1","Hapus akun → e-mail bebas dipakai daftar ulang, sesi lama mati","Akun uji","S35","admin_delete_partner('user') lalu daftar ulang e-mail sama","Tombstone auth.users+identities; daftar ulang berhasil","Sesuai (S35a–e)","-","bukti/simulasi_e2e.log","PASS",""),
 ("AUTH-05","6.1","Lupa/reset kata sandi lewat e-mail, tautan kedaluwarsa","Supabase Auth e-mail","-","Alur forgot → kode → reset","Berhasil; kode salah/kedaluwarsa ditolak","Bukti sesi 6 Sep (19 pemeriksaan) — TIDAK diulang pada RC ini","-","docs/LUPA-KATA-SANDI.md","BLOCKED","QA: ulang pada RC final"),
 ("AUTH-06","6.1 / 6.7","Proteksi kata sandi bocor (HIBP) & pembatasan percobaan login","Dashboard Supabase","-","Cek Auth → Attack Protection","Leaked Password Protection ON","Advisor Supabase: auth_leaked_password_protection = WARN (OFF)","Medium","get_advisors security 12 Sep","FAIL","Erza (login dashboard) — sebelum closed beta"),
 # 6.2 Order & state machine
 ("ORD-01","6.2 / E2E-01","Golden path pesanan: searching → accepted → arrived → in_progress → completed, PIN serah-terima","Akun uji customer+driver, saldo","ride_motor Rp9.000 dompet","S1 + S38a","Status berurutan; PIN benar; dompet: pelanggan −total, driver +bagian","Sesuai","-","bukti/simulasi_e2e.log S1, S38","PASS",""),
 ("ORD-02","6.2","Satu order tidak dapat diterima dua mitra (balapan)","2 driver uji","S48a","Dua driver_accept_order berturutan","Satu pemenang, satu ditolak, event accepted=1","Sesuai","-","bukti/simulasi_e2e.log S48a–b","PASS",""),
 ("ORD-03","6.2 / E2E-07","Ketuk ganda / koneksi putus saat submit → tidak ada order & tagihan ganda","Akun uji customer","create_order 2× identik; 2× dengan client_request_id sama","Panggil create_order dua kali","1 order, 1 potongan dompet","SEBELUM: 2 order + 2 potongan (S48c). SESUDAH migrasi 0083 + orders.ts: 1 order, 1 potongan","High (diperbaiki)","bukti/uji_idempotensi.log; scripts/test-order-request.mjs 7/7","PASS (setelah perbaikan)","Build 105 memuat kunci klien; server sudah aktif"),
 ("ORD-04","6.2 / E2E-06","Tidak ada mitra / mitra menolak / kelas kendaraan tidak tersedia","Driver uji","S17, S23, S24","Tolak order; tunggu masa tahan; prioritas rating","Order hilang dari feed penolak; fallback kelas; prioritas rating; masa tahan AntarNow","Sesuai","-","bukti/simulasi_e2e.log S17, S22–S24, S34","PASS",""),
 ("ORD-05","6.2 / E2E-05","Pembatalan pada tiap status yang diizinkan, konsekuensi biaya/refund","Akun uji","S44a–g, S50c, S50e, S51e","Batalkan di SEARCHING/ACCEPTED/ARRIVED/IN_PROGRESS/scheduled; merchant menolak","Refund utuh kecuali kebijakan (carter <12 jam potong 30%); IN_PROGRESS tidak bisa dibatalkan sendiri","Sesuai","-","bukti/simulasi_e2e.log S44, S45, S50, S51","PASS",""),
 ("ORD-06","6.2 / E2E-08","Intervensi admin: batal/refund dengan alasan, pelaku, jejak audit, PIN","Akun admin uji","S45, S27","Admin batalkan IN_PROGRESS; hapus mitra tanpa PIN/alasan","cancelled_by=admin, refund utuh; tanpa PIN/alasan ditolak; security_events bertambah","Sesuai","-","bukti/simulasi_e2e.log S27, S45","PASS",""),
 ("ORD-07","6.2","Mitra tidak bisa menyelesaikan/mengambil order di luar assignment","Driver uji","S23c, S48","Tolak/terima order yang sudah dimiliki driver lain","Ditolak","Sesuai","-","bukti/simulasi_e2e.log S23c, S48a","PASS",""),
 # 6.3 Uang
 ("PAY-01","6.3","Quote harga dapat dilacak: jarak, tarif, biaya jasa, promo, total; bagian driver/merchant/platform","-","S29d admin_order_split","Bandingkan split order vs laporan","Identitas revenue−promo−gateway=net; cogs=driver+merchant+gateway+promo","Sesuai (0 pelanggaran)","-","bukti/simulasi_e2e.log S29, S30","PASS",""),
 ("PAY-02","6.3","Klien tidak bisa memanipulasi nominal: server menghitung ulang dari route_km; tagihan = yang ditampilkan","-","tests/peta/uji-ongkos.mjs","Pratinjau vs finalizeRoute untuk 7 layanan","route_km dikirim = jarak pratinjau dasar harga","Sesuai 7/7 layanan","-","tests/peta/uji-ongkos.mjs (12 Sep)","PASS",""),
 ("PAY-03","6.3 / E2E-03","Webhook settlement dikirim 2×, status expire, external_id tak dikenal","Payment uji","S46c–e","payment_settle 2× ref sama","Kredit sekali; expire Δ0; tak dikenal ditolak","Sesuai (kredit ke-2 = 0, 1 baris mutasi)","-","bukti/simulasi_e2e.log S46","PASS",""),
 ("PAY-04","6.3","Webhook Midtrans dengan signature palsu / payload kosong","Edge Function midtrans-webhook (verify_jwt=false)","signature_key='deadbeef'","POST langsung dari browser","Ditolak, tidak menyentuh saldo","400 'Server key belum diatur' (gagal-tertutup)","-","uji langsung 12 Sep (Chrome)","PASS","Ulangi dengan kunci sandbox terpasang"),
 ("PAY-05","6.3","Snap sandbox: sukses, gagal, pending, kedaluwarsa, dibayar terlambat","Kunci Midtrans sandbox di gateway_secrets + Notification URL","-","Top up dari aplikasi → bayar di Snap sandbox","Saldo bertambah otomatis via webhook; kasus gagal tidak menambah saldo","gateway_secrets kosong — belum bisa diuji","-","-","BLOCKED","Erza: login dashboard Midtrans (sandbox + production)"),
 ("PAY-06","6.3","Refund/penyesuaian saldo manual admin, top up manual ditolak","Admin uji","S45c, S46a–b","Adjust +25.000; tolak top up; setujui ulang","Mutasi adjustment 1 baris; ditolak Δ0; persetujuan ulang ditolak","Sesuai","-","bukti/simulasi_e2e.log S45c, S46a–b","PASS",""),
 ("PAY-07","6.3 (rekonsiliasi)","Invarian buku besar: saldo = Σ mutasi; 0 order selesai tanpa pembagian; 0 dompet minus non-driver; platform tidak membayar > diskon","Seluruh data simulasi","S53a–e","Hitung invarian di akhir simulasi","0 pelanggaran","0 pelanggaran (5/5)","-","bukti/simulasi_e2e.log S53","PASS",""),
 ("PAY-08","6.3","Tidak ada kunci gateway/API di repo, log, aplikasi","Repo","git grep pola kunci","Pindai repo (kecuali lockfile)","0 temuan","0 temuan (hanya placeholder dokumentasi)","-","git grep 12 Sep","PASS",""),
 ("PAY-09","6.3","Penarikan: minimum, negatif, melebihi saldo, ganda, ditolak admin, otomatis vs manual saat flag fraud","Driver uji","S12, S47","7 skenario penarikan","Semua ditolak/diterima sesuai aturan; saldo tak pernah minus","Sesuai","-","bukti/simulasi_e2e.log S12, S47","PASS",""),
 ("PAY-10","6.3 / hukum","Komisi roda dua ≤ 8% (Perpres 27/2026) dijaga server","Pricing","S37","Set komisi 9%","Ditolak pengawal","Ditolak; 7% diterima; non-roda-dua tak terpengaruh","-","bukti/simulasi_e2e.log S37","PASS","Food/send: menunggu opini hukum"),
 ("PAY-11","6.3","Order tunai: potongan platform dari dompet driver, deposit minus dibatasi −500.000","Driver uji","S39–S42, S52","Order tunai 5 layanan + driver saldo rendah","Sisa di tangan driver = pendapatan; di bawah ambang ditolak menerima order","Sesuai","-","bukti/simulasi_e2e.log S39–S42, S52","PASS",""),
 # 6.4 Lokasi
 ("LOC-01","6.4","Gerbang kota di SERVER (bukan hanya UI): kota belum dilayani / di luar radius ditolak dengan alasan","Data kota","uji_kota.sql + Playwright","Pesan dari Medan; titik 961 km; lewati UI panggil RPC","Ditolak dengan pesan Bahasa Indonesia; UI konsisten","29 lulus (Playwright) + 12/12 blok SQL","-","bukti/uji_kota.log; tests/kota/uji-kota.mjs","PASS",""),
 ("LOC-02","6.4","Batas jarak per layanan & antar kota","-","S18a–g","Motor 262 km; food 20 km; send antar kota","Ditolak dengan pesan yang mengarahkan ke layanan lain","Sesuai","-","bukti/simulasi_e2e.log S18","PASS",""),
 ("LOC-03","6.4 / 6.7","Kunci peta: rahasia tidak pernah keluar ke klien; kunci publik disisipkan server; Stadia aktif","map_secrets","uji_peta.sql; uji langsung Stadia","map_public_config() sebagai anon; fetch ubin/geocode/rute","Tanpa secret_key; 200 untuk 4 endpoint Stadia","Sesuai; uses_free_osm=false","-","bukti/uji_peta.log; uji langsung 12 Sep","PASS","Kunci Stadia tidak bisa dibatasi per package Android — pantau kuota"),
 ("LOC-04","6.4","Izin lokasi ditolak → aplikasi tetap memberi jalur aman (pilih titik manual)","Perangkat Android","-","Tolak izin lokasi lalu pesan","Place-picker manual tetap bisa","Belum diuji pada perangkat fisik pada RC ini","-","-","BLOCKED","QA manual di HP (Erza/tester) — sebelum closed beta"),
 ("LOC-05","6.4","Tombol bantuan/SOS & tiket eskalasi","-","S13, S28","Buat tiket, SOS, tutup dengan rating","Tiket tercatat, SOS ditangani, dedupe tiket","Sesuai","-","bukti/simulasi_e2e.log S13, S28","PASS","SOP respons SOS manusia belum diuji lapangan"),
 # 6.5 Merchant & mitra
 ("MER-01","6.5","Merchant tutup/belum disetujui tidak menerima order; keranjang kosong ditolak","Merchant uji","create_order food","Pesan ke merchant tutup","Ditolak","Ditolak (aturan create_order; S3 jalur sukses)","-","supabase/migrations/0014 + S3","PASS",""),
 ("MER-02","6.5","Merchant menolak order berbayar → refund & kuota promo kembali","-","S44g","Merchant reject","Refund utuh, used_count kembali","Sesuai","-","bukti/simulasi_e2e.log S44g","PASS",""),
 ("MER-03","6.5","Pendapatan, insentif, potongan mitra normal & pembatalan","-","S38, S44, S50","Bandingkan Δ dompet","Sesuai rumus tiap layanan","Sesuai","-","bukti/simulasi_e2e.log S38, S44, S50","PASS",""),
 ("MER-04","6.5","Dokumen mitra: akses minimum & retensi","-","-","Review RLS dokumen mitra + kebijakan retensi","Hanya pemilik & admin; retensi tertulis","RLS ada; kebijakan retensi belum ditinjau legal","-","-","BLOCKED","Review privacy/legal (bukan QA)"),
 # 6.6 Notifikasi
 ("NOT-01","6.6","Notifikasi membawa kunci tujuan; tidak mengungkap data sensitif","-","S33a–g","Periksa payload notifikasi","order_id/ticket_id/booking_id ada; isi netral","Sesuai; 5 jenis informatif tanpa tujuan (dicatat)","-","bukti/simulasi_e2e.log S33","PASS",""),
 ("NOT-02","6.6","Token push idempoten & RLS; pengiriman ulang tidak mengubah status","-","S36","register_push_token 2×; pengguna lain mencabut","1 baris; pengguna lain tidak bisa","Sesuai","-","bukti/simulasi_e2e.log S36","PASS",""),
 ("NOT-03","6.6","Push benar-benar sampai ke perangkat (FCM v1) dari server","Secret FCM_SERVICE_ACCOUNT + push_config","-","Kirim notifikasi → HP","Notifikasi tampil","Secret belum diisi → push_dispatch skipped","-","S36f","BLOCKED","Erza: unggah service account Firebase ke Supabase secrets + admin_set_push_config"),
 ("NOT-04","6.6","Moderasi UGC: blokir pengguna, laporan konten, batas laju, jejak audit","-","uji_moderasi.sql","18 pemeriksaan","Semua lulus","18/18","-","bukti/uji_moderasi.log","PASS",""),
 # 6.7 Keamanan & ketahanan
 ("SEC-01","6.7","TLS, CORS, error handling Edge Function","-","-","Periksa header & respons","HTTPS saja; error tanpa detail internal","Supabase TLS; CORS '*' + JWT; error pesan ringkas","-","kode Edge Functions","PASS","CORS '*' dapat diterima untuk API publik ber-JWT"),
 ("SEC-02","6.7","Dependency scan (npm audit)","-","npm audit 12 Sep","-","0 high/critical","0 high/critical; 15 moderate transitif dari 2 advisori (decode-uri-component via expo-router; uuid via xcode build-time)","Medium","npm audit 12 Sep","PASS (GO WITH RISK)","Eng: naikkan saat Expo SDK berikutnya; tidak ada perbaikan non-breaking"),
 ("SEC-03","6.7","Secret scan repo","-","git grep","-","0 rahasia","0 rahasia","-","git grep 12 Sep","PASS",""),
 ("SEC-04","6.7 / E2E-10","Backup/restore diuji (bukti restore DB & pemulihan order/payment)","Backup Supabase","-","Latihan restore ke proyek latihan","Rekonsiliasi kembali dapat dijelaskan","Belum pernah dilakukan; paket Supabase saat ini belum diverifikasi punya PITR","-","-","BLOCKED","Eng+Erza: aktifkan backup harian/PITR, latihan restore sebelum rilis nasional"),
 ("SEC-05","6.7","SAST/DAST/penetration test independen jalur kritis","-","-","-","Tidak ada finding High/Critical tanpa mitigasi","Belum dilakukan pihak independen (baru advisor Supabase + uji internal)","-","get_advisors 12 Sep","BLOCKED","Erza: vendor pentest sebelum rilis nasional"),
 ("SEC-06","6.7","Runbook insiden, rollback, on-call, status page","-","docs/rilis/*","-","Tersedia & pernah dilatih","Runbook rilis ada; latihan insiden & status page belum","-","docs/rilis/RUNBOOK-LISTING-BESOK.md","PASS (GO WITH RISK)","Ops: latihan insiden 1× sebelum closed beta"),
 ("SEC-07","6.7","Advisor keamanan Supabase (RLS, security definer, extension in public)","-","get_advisors","-","Tidak ada ERROR tak terjelaskan","1 ERROR: spatial_ref_sys (tabel sistem PostGIS, tidak berisi data pengguna); 23 fungsi anon = jalur publik pra-login yang disengaja (estimasi, status kota, shared_order)","Low","get_advisors 12 Sep","PASS (dengan catatan)","Eng: dokumentasikan daftar fungsi anon yang disengaja"),
 # 6.8 Performa
 ("PERF-01","6.8","Uji perangkat/versi Android target, layar kecil, memori rendah, jaringan putus-sambung","HP fisik","-","Sesi eksplorasi terstruktur","Tidak ada freeze/crash di checkout & pelacakan","Belum dilakukan pada RC ini","-","-","BLOCKED","QA manual (Erza + 2 tester) sebelum closed beta; Play pre-launch report"),
 ("PERF-02","6.8","SLO & uji beban (lonjakan order, notifikasi serentak, webhook ulang)","-","-","-","Kapasitas & bottleneck tercatat","Belum dilakukan","-","-","BLOCKED","Eng: k6/ artillery terhadap RPC create_order & webhook — sebelum rilis nasional"),
 ("PERF-03","6.8","Aksesibilitas: TalkBack, teks besar, kontras, label kontrol","-","-","-","Lolos pemeriksaan dasar","Belum dilakukan","-","-","BLOCKED","QA manual + Play pre-launch accessibility"),
 # Infrastruktur baru
 ("CALL-01","6.6 / 6.8","Panggilan suara di jaringan seluler: kredensial TURN berumur pendek, api_token tidak bocor","Cloudflare TURN","uji_turn.sql + uji langsung","anon → fungsi; admin → fungsi; RPC admin","anon 401; admin dapat iceServers; 9/9 SQL","Sesuai","-","bukti/uji_turn.log","PASS","Uji panggilan nyata 2 HP di jaringan seluler: BLOCKED (manual)"),
 ("DATA-01","6.4 (data tempat)","Impor tempat OSM: gazetteer, dedupe, faskes ke poi, kota layanan tidak tergandakan","-","uji_impor.sql","12 skenario","Semua sesuai","Sesuai; S9 terhalang pekerjaan impor yang sedang berjalan (efek data)","Low","bukti/uji_impor.log","PASS (dengan catatan)","Pekerjaan impor nasional: 7/39 tugas; Overpass publik lambat — 0084 + mirror ketiga dipasang"),
]

E2E = [
 ("E2E-01","Antar barang bayar dompet → mitra → pickup → antar → selesai (QRIS diganti dompet karena Midtrans belum terpasang)","PASS (server) — ORD-01, PAY-07; jalur QRIS BLOCKED (PAY-05)"),
 ("E2E-02","Antar makan: merchant terima, mitra pickup, antar, selesai; harga sama di semua peran","PASS — S3, S38c, S42, PAY-01"),
 ("E2E-03","Webhook pembayaran sama dikirim 2× → satu pembayaran","PASS — S46c (DB) + PAY-04 (HTTP gagal-tertutup); ulang dengan kunci sandbox"),
 ("E2E-04","Pembayaran sukses tetapi mitra tak tersedia → refund/batal, saldo tidak menggantung","PASS — S44a (batal saat SEARCHING refund utuh), S51e"),
 ("E2E-05","Batal pada tiap titik status","PASS — ORD-05"),
 ("E2E-06","Mitra offline saat ditugaskan/mengantar → timeout/redispatch","PASS sebagian — S17/S23/S24/S34 (fallback & masa tahan); skenario 'offline saat mengantar' hanya lewat admin batal (S45a): perlu uji lapangan"),
 ("E2E-07","Koneksi putus saat submit → tidak ada order/tagihan ganda","PASS setelah perbaikan 0083 — ORD-03"),
 ("E2E-08","Admin menyelesaikan dispute/refund dengan otorisasi & jejak","PASS — ORD-06, PAY-06"),
 ("E2E-09","Pengguna A akses order pengguna B","PASS — AUTH-01, S31c–f"),
 ("E2E-10","Restore dari backup di lingkungan latihan","BLOCKED — SEC-04"),
]

def build_xlsx(path):
    wb = Workbook(); ws = wb.active; ws.title = "Matriks Uji"
    hdr = ["ID","Bagian standar","Skenario","Prasyarat","Data uji","Langkah","Expected","Actual","Severity","Bukti","Status","PIC / target"]
    f = Font(name="Arial", size=10); fb = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0F766E"); thin = Side(style="thin", color="D0D7DE"); bd = Border(left=thin,right=thin,top=thin,bottom=thin)
    ws.append([f"Matriks Uji AntarKita — Standar Testing Listing Nasional v1.0 · {TGL}"]); ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws.append([f"Release candidate: {RC}"]); ws["A2"].font = f
    ws.append([f"Pelaksana: {PEL} · Lingkungan: basis data produksi (setiap uji di-ROLLBACK), web live, Edge Functions"]); ws["A3"].font = f
    ws.append([]); ws.append(hdr)
    for c in ws[5]: c.font = fb; c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = bd
    colors = {"PASS":"DCFCE7","FAIL":"FEE2E2","BLOCKED":"FEF3C7"}
    for row in T:
        ws.append(list(row))
        r = ws.max_row
        for c in ws[r]: c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = bd
        st = row[10]; key = "PASS" if st.startswith("PASS") else ("FAIL" if st.startswith("FAIL") else "BLOCKED")
        ws.cell(r, 11).fill = PatternFill("solid", fgColor=colors[key])
    widths = [10,16,44,22,22,30,32,40,12,30,18,34]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"
    # Ringkasan dengan rumus
    s = wb.create_sheet("Ringkasan")
    s["A1"] = "Ringkasan status (dihitung dari sheet Matriks Uji)"; s["A1"].font = Font(name="Arial", size=12, bold=True)
    n0, n1 = 6, 5 + len(T)
    rows = [("PASS (termasuk 'PASS (…)')", f'=COUNTIF(\'Matriks Uji\'!K{n0}:K{n1},"PASS*")'),
            ("FAIL", f'=COUNTIF(\'Matriks Uji\'!K{n0}:K{n1},"FAIL*")'),
            ("BLOCKED", f'=COUNTIF(\'Matriks Uji\'!K{n0}:K{n1},"BLOCKED*")'),
            ("Total kasus uji", f"=COUNTA('Matriks Uji'!A{n0}:A{n1})"),
            ("Severity High (diperbaiki)", f'=COUNTIF(\'Matriks Uji\'!I{n0}:I{n1},"High*")'),
            ("Severity Medium terbuka", f'=COUNTIF(\'Matriks Uji\'!I{n0}:I{n1},"Medium")')]
    for i,(k,v) in enumerate(rows, start=3):
        s.cell(i,1,k).font = f; s.cell(i,2,v).font = f
    s["A10"] = "Keputusan"; s["A10"].font = Font(name="Arial", size=10, bold=True)
    s["B10"] = "INTERNAL TEST: GO WITH RISK · NATIONAL LISTING READY: NO-GO (lihat laporan)"; s["B10"].font = f
    s["A12"] = "Catatan keterbatasan: uji server dijalankan pada basis data produksi dalam transaksi ROLLBACK; uji perangkat fisik, beban, aksesibilitas, pentest, dan restore backup belum dilakukan pada RC ini."; s["A12"].font = f
    s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 90
    e = wb.create_sheet("E2E wajib"); e.append(["ID","Skenario (Standar §7)","Hasil / rujukan"])
    for c in e[1]: c.font = fb; c.fill = fill
    for row in E2E:
        e.append(list(row))
        for c in e[e.max_row]: c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
    e.column_dimensions["A"].width = 10; e.column_dimensions["B"].width = 70; e.column_dimensions["C"].width = 80
    wb.save(path)

def md_tables():
    out = ["| ID | Bagian | Skenario | Actual | Severity | Bukti | Status | PIC / target |", "|---|---|---|---|---|---|---|---|"]
    for r in T:
        out.append(f"| {r[0]} | {r[1]} | {r[2]} | {r[7]} | {r[8]} | {r[9]} | **{r[10]}** | {r[11]} |")
    out2 = ["| ID | Skenario | Hasil |", "|---|---|---|"] + [f"| {a} | {b} | {c} |" for a,b,c in E2E]
    return "\n".join(out), "\n".join(out2)

if __name__ == "__main__":
    import sys, os
    here = os.path.dirname(os.path.abspath(__file__))
    build_xlsx(os.path.join(here, "Matriks-Uji-AntarKita-2026-09-12.xlsx"))
    a, b = md_tables()
    open(os.path.join(here, "_tabel-matriks.md"), "w").write(a + "\n\n" + b + "\n")
    print("ok", len(T), "kasus,", len(E2E), "E2E")
